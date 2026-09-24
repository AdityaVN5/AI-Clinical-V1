"""JSON, PDF and DOCX export helpers."""

from __future__ import annotations

import io
import json

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from docx import Document


DISCLAIMER = (
    "AI-generated content is intended for documentation assistance only. "
    "It must be reviewed by a qualified healthcare professional before being "
    "used in a medical record or for clinical decision-making."
)


def export_json(record: dict) -> bytes:
    """Export the EHR record as formatted JSON."""

    return json.dumps(
        record,
        indent=2,
        ensure_ascii=False,
        default=str,
    ).encode("utf-8")


def _safe(value) -> str:
    """Convert empty or missing values into a readable string."""

    if value is None:
        return "Not documented."

    if isinstance(value, str) and not value.strip():
        return "Not documented."

    return str(value)


def _pdf_text(value) -> str:
    """Prepare text safely for ReportLab Paragraph."""

    text = _safe(value)

    # Escape basic XML/HTML characters that could break ReportLab Paragraph.
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    # Preserve line breaks.
    text = text.replace("\n", "<br/>")

    return text


def export_pdf(record: dict) -> bytes:
    """
    Export the EHR record as a PDF.

    Includes:
    - Patient information
    - Consultation transcript
    - SOAP note
    - Compliance result
    - ICD-10 suggestions
    - Healthcare disclaimer
    """

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph(
            "AI Clinical Documentation Assistant",
            styles["Title"],
        ),
        Spacer(1, 12),
    ]

    # =========================================================
    # PATIENT INFORMATION
    # =========================================================

    patient = record.get("patient", {})

    story.append(
        Paragraph(
            "Patient Information",
            styles["Heading2"],
        )
    )

    patient_text = (
        f"<b>Patient ID:</b> {_pdf_text(patient.get('patient_id'))}<br/>"
        f"<b>Patient Name:</b> {_pdf_text(patient.get('patient_name'))}<br/>"
        f"<b>Age:</b> {_pdf_text(patient.get('age'))}<br/>"
        f"<b>Gender:</b> {_pdf_text(patient.get('gender'))}"
    )

    story.append(
        Paragraph(
            patient_text,
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 12))

    # =========================================================
    # CONSULTATION TRANSCRIPT
    # =========================================================

    consultation = record.get("consultation", {})

    story.append(
        Paragraph(
            "Consultation Transcript",
            styles["Heading2"],
        )
    )

    transcript = _pdf_text(
        consultation.get("transcript")
    )

    story.append(
        Paragraph(
            transcript,
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 12))

    # =========================================================
    # SOAP NOTE
    # =========================================================

    story.append(
        Paragraph(
            "SOAP Note",
            styles["Heading2"],
        )
    )

    soap_sections = [
        ("Subjective", "subjective"),
        ("Objective", "objective"),
        ("Assessment", "assessment"),
        ("Plan", "plan"),
    ]

    for title, key in soap_sections:

        story.append(
            Paragraph(
                title,
                styles["Heading3"],
            )
        )

        story.append(
            Paragraph(
                _pdf_text(consultation.get(key)),
                styles["BodyText"],
            )
        )

        story.append(Spacer(1, 6))

    story.append(Spacer(1, 6))

    # =========================================================
    # COMPLIANCE
    # =========================================================

    compliance = record.get("compliance", {})

    story.append(
        Paragraph(
            "Compliance",
            styles["Heading2"],
        )
    )

    compliance_score = compliance.get(
        "score",
        compliance.get("compliance_score", 0),
    )

    compliance_status = compliance.get(
        "status",
        "Incomplete",
    )

    story.append(
        Paragraph(
            (
                f"<b>Documentation Completeness Score:</b> "
                f"{_pdf_text(compliance_score)}/100<br/>"
                f"<b>Status:</b> {_pdf_text(compliance_status)}"
            ),
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 8))

    # Passed checks
    passed_checks = compliance.get(
        "passed_checks",
        [],
    )

    story.append(
        Paragraph(
            "Passed Checks",
            styles["Heading3"],
        )
    )

    if passed_checks:

        for item in passed_checks:

            if item:
                story.append(
                    Paragraph(
                        f"• {_pdf_text(item)}",
                        styles["BodyText"],
                    )
                )

    else:

        story.append(
            Paragraph(
                "None.",
                styles["BodyText"],
            )
        )

    story.append(Spacer(1, 6))

    # Missing information
    missing_information = compliance.get(
        "missing_information",
        [],
    )

    story.append(
        Paragraph(
            "Missing Information",
            styles["Heading3"],
        )
    )

    if missing_information:

        for item in missing_information:

            if item:
                story.append(
                    Paragraph(
                        f"• {_pdf_text(item)}",
                        styles["BodyText"],
                    )
                )

    else:

        story.append(
            Paragraph(
                "None.",
                styles["BodyText"],
            )
        )

    story.append(Spacer(1, 6))

    # Warnings
    warnings = compliance.get(
        "warnings",
        [],
    )

    story.append(
        Paragraph(
            "Warnings",
            styles["Heading3"],
        )
    )

    if warnings:

        for item in warnings:

            if item:
                story.append(
                    Paragraph(
                        f"• {_pdf_text(item)}",
                        styles["BodyText"],
                    )
                )

    else:

        story.append(
            Paragraph(
                "None.",
                styles["BodyText"],
            )
        )

    story.append(Spacer(1, 6))

    # Recommendations
    recommendations = compliance.get(
        "recommendations",
        [],
    )

    story.append(
        Paragraph(
            "Recommendations",
            styles["Heading3"],
        )
    )

    if recommendations:

        for item in recommendations:

            if item:
                story.append(
                    Paragraph(
                        f"• {_pdf_text(item)}",
                        styles["BodyText"],
                    )
                )

    else:

        story.append(
            Paragraph(
                "None.",
                styles["BodyText"],
            )
        )

    story.append(Spacer(1, 12))

    # =========================================================
    # ICD-10 SUGGESTIONS
    # =========================================================

    story.append(
        Paragraph(
            "ICD-10 Suggestions",
            styles["Heading2"],
        )
    )

    icd_codes = record.get(
        "icd_codes",
        [],
    )

    # Make sure ICD data is a list.
    if not isinstance(icd_codes, list):
        icd_codes = []

    data = [
        [
            "Code",
            "Description",
            "Reason",
        ]
    ]

    for item in icd_codes:

        if not isinstance(item, dict):
            continue

        data.append(
            [
                _safe(item.get("code")),
                _safe(item.get("description")),
                _safe(item.get("reason")),
            ]
        )

    if len(data) == 1:

        data.append(
            [
                "—",
                "No reliable suggestion",
                "Clinician/coder verification required.",
            ]
        )

    # =========================================================
    # ICD TABLE
    # =========================================================

    table = Table(
        data,
        colWidths=[
            65,
            150,
            300,
        ],
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                # IMPORTANT:
                # ReportLab requires the line color as the 5th
                # argument for GRID.
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.black,
                ),

                # Header background
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                # Header font
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                # Vertical alignment
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),

                # Padding
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(table)

    story.append(Spacer(1, 12))

    # =========================================================
    # DISCLAIMER
    # =========================================================

    story.append(
        Paragraph(
            DISCLAIMER,
            styles["Italic"],
        )
    )

    # =========================================================
    # BUILD PDF
    # =========================================================

    document.build(story)

    return buffer.getvalue()


def export_docx(record: dict) -> bytes:
    """
    Export the EHR record as a DOCX document.
    """

    document = Document()

    # =========================================================
    # TITLE
    # =========================================================

    document.add_heading(
        "AI Clinical Documentation Assistant",
        level=0,
    )

    # =========================================================
    # PATIENT INFORMATION
    # =========================================================

    patient = record.get(
        "patient",
        {},
    )

    document.add_heading(
        "Patient Information",
        level=1,
    )

    document.add_paragraph(
        f"Patient ID: {_safe(patient.get('patient_id'))}\n"
        f"Patient Name: {_safe(patient.get('patient_name'))}\n"
        f"Age: {_safe(patient.get('age'))}\n"
        f"Gender: {_safe(patient.get('gender'))}"
    )

    # =========================================================
    # CONSULTATION
    # =========================================================

    consultation = record.get(
        "consultation",
        {},
    )

    document.add_heading(
        "Consultation Transcript",
        level=1,
    )

    document.add_paragraph(
        _safe(
            consultation.get("transcript")
        )
    )

    # =========================================================
    # SOAP NOTE
    # =========================================================

    document.add_heading(
        "SOAP Note",
        level=1,
    )

    for title, key in [
        ("Subjective", "subjective"),
        ("Objective", "objective"),
        ("Assessment", "assessment"),
        ("Plan", "plan"),
    ]:

        document.add_heading(
            title,
            level=2,
        )

        document.add_paragraph(
            _safe(
                consultation.get(key)
            )
        )

    # =========================================================
    # COMPLIANCE
    # =========================================================

    compliance = record.get(
        "compliance",
        {},
    )

    document.add_heading(
        "Documentation Completeness",
        level=1,
    )

    score = compliance.get(
        "score",
        compliance.get("compliance_score", 0),
    )

    status = compliance.get(
        "status",
        "Incomplete",
    )

    document.add_paragraph(
        f"{score}/100 — {status}"
    )

    # Missing information
    document.add_heading(
        "Missing Information",
        level=2,
    )

    missing = compliance.get(
        "missing_information",
        [],
    )

    if missing:

        for item in missing:

            if item:
                document.add_paragraph(
                    str(item),
                    style="List Bullet",
                )

    else:

        document.add_paragraph(
            "None."
        )

    # Warnings
    document.add_heading(
        "Warnings",
        level=2,
    )

    warnings = compliance.get(
        "warnings",
        [],
    )

    if warnings:

        for item in warnings:

            if item:
                document.add_paragraph(
                    str(item),
                    style="List Bullet",
                )

    else:

        document.add_paragraph(
            "None."
        )

    # Recommendations
    document.add_heading(
        "Recommendations",
        level=2,
    )

    recommendations = compliance.get(
        "recommendations",
        [],
    )

    if recommendations:

        for item in recommendations:

            if item:
                document.add_paragraph(
                    str(item),
                    style="List Bullet",
                )

    else:

        document.add_paragraph(
            "None."
        )

    # =========================================================
    # ICD-10
    # =========================================================

    document.add_heading(
        "ICD-10 Suggestions",
        level=1,
    )

    icd_codes = record.get(
        "icd_codes",
        [],
    )

    if isinstance(icd_codes, list) and icd_codes:

        valid_codes = False

        for item in icd_codes:

            if not isinstance(item, dict):
                continue

            valid_codes = True

            code = _safe(
                item.get("code")
            )

            description = _safe(
                item.get("description")
            )

            reason = _safe(
                item.get("reason")
            )

            document.add_paragraph(
                (
                    f"{code} — {description}\n"
                    f"Reason: {reason}"
                ),
                style="List Bullet",
            )

        if not valid_codes:
            document.add_paragraph(
                "No reliable suggestion."
            )

    else:

        document.add_paragraph(
            "No reliable suggestion."
        )

    # =========================================================
    # DISCLAIMER
    # =========================================================

    document.add_heading(
        "Disclaimer",
        level=1,
    )

    document.add_paragraph(
        DISCLAIMER
    )

    # =========================================================
    # SAVE DOCX TO MEMORY
    # =========================================================

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()